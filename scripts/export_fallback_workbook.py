from __future__ import annotations

import argparse
import json
import os
import sys
import time
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

REPO_ROOT = Path(__file__).resolve().parents[1]
GWS_CONFIG_DIR = str(Path.home() / ".config" / "gws-amazon")
REGIONS = ["us", "de", "uk", "fr", "es"]


def _xlsx_sheet_title(title: str) -> str:
    """Excel forbids []:*?/\\ in worksheet titles; keep the source mapping in manifest."""
    return title.translate(str.maketrans("", "", "[]:*?/\\"))[:31]


def _env() -> None:
    os.environ.setdefault("AMZ_BS_ROOT_BACKEND", "gws")
    os.environ.setdefault("AMZ_BS_ROOT_SHEETS_ID", "1XQoI7SSuFKbuRAeD23uQIfJu3FBXEv__6rvm68Zfo80")
    os.environ.setdefault("GOOGLE_WORKSPACE_CLI_CONFIG_DIR", GWS_CONFIG_DIR)
    os.environ.setdefault("GOOGLE_WORKSPACE_CLI_KEYRING_BACKEND", "keyring")


def _write_grid(wb: Workbook, title: str, grid: list[list[str]], stamp: list[str] | None = None) -> None:
    ws = wb.create_sheet(_xlsx_sheet_title(title))
    if stamp:
        ws.append(stamp + [""] * max(0, len(grid[0]) - len(stamp)))
    for row in grid:
        ws.append(list(row))
    header_row = 2 if stamp else 1
    for cell in ws[header_row]:
        cell.font = Font(bold=True)
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for cells in ws.columns:
        width = max((len(str(c.value or "")) for c in cells), default=8)
        ws.column_dimensions[get_column_letter(cells[0].column)].width = min(max(width + 2, 10), 55)
    ws.freeze_panes = f"A{header_row + 1}"


def build_workbook(output: Path) -> None:
    _env()
    sys.path.insert(0, str(REPO_ROOT / "src"))
    from crawl_amazon_beauty_bestsellers.config import load_settings
    from crawl_amazon_beauty_bestsellers.root_publish import (
        all_top100_grid,
        cross_region_catalog_grid,
        price_history_grid,
        rank_changes_grid,
        rank_history_pivot_grid,
        root_history_rows,
        root_panel_grid,
        root_trend_grid,
    )
    from crawl_amazon_beauty_bestsellers.storage.store import Store

    settings = load_settings(REPO_ROOT)
    store = Store(settings)
    try:
        wb = Workbook()
        wb.remove(wb.active)
        generated = time.strftime("%Y-%m-%d %H:%M %Z")
        stamp = [generated, "source=local SQLite", "fallback workbook"]
        for region in REGIONS:
            _write_grid(wb, f"[{region.upper()}] Top 100", root_panel_grid(store, settings, region), stamp)
        _write_grid(wb, "[ALL] Top 100", all_top100_grid(store, settings), stamp)

        history_rows: list[list[str]] = []
        for region in REGIONS:
            history_rows.extend(root_history_rows(store, settings, region))
        history_rows.sort(key=lambda row: (row[0], row[1], int(row[3]) if row[3].isdigit() else 999))
        _write_grid(
            wb,
            "root_rank_history",
            [["date", "region", "asin", "rank", "price_krw", "currency", "rating", "ratings_count", "title"], *history_rows],
        )
        for region in REGIONS:
            _write_grid(wb, f"{region.upper()} Rank History", rank_history_pivot_grid(store, settings, region), stamp)
        _write_grid(wb, "trend_14d", root_trend_grid(store, settings, "es"), stamp)
        _write_grid(wb, "Cross-Region Catalog", cross_region_catalog_grid(store, settings), stamp)
        _write_grid(wb, "Rank Changes", rank_changes_grid(store, settings, "es"), stamp)
        _write_grid(wb, "Price History", price_history_grid(store, settings, "es"), stamp)

        manifest = wb.create_sheet("_fallback_manifest")
        manifest.append(["generated_at", generated])
        manifest.append(["source", "local SQLite; independent of Google Sheets read/write availability"])
        manifest.append(["purpose", "daily recovery artifact and email fallback"])
        manifest.append(["source_sheet_tab", "xlsx_tab"])
        for region in REGIONS:
            source_tab = f"[{region.upper()}] Top 100"
            manifest.append([source_tab, _xlsx_sheet_title(source_tab)])
        manifest.append(["[ALL] Top 100", _xlsx_sheet_title("[ALL] Top 100")])
        manifest["A1"].font = Font(bold=True)
        output.parent.mkdir(parents=True, exist_ok=True)
        wb.save(output)
    finally:
        store.close()


def main() -> int:
    parser = argparse.ArgumentParser(description="Create an independent local-data XLSX fallback workbook.")
    parser.add_argument("--email", action="store_true")
    parser.add_argument("--no-upload", action="store_true")
    parser.add_argument("--summary-json", type=Path, help="Write the fallback result for an enclosing report.")
    args = parser.parse_args()
    day = time.strftime("%Y%m%d")
    output = REPO_ROOT / "artifacts" / "exports" / "xlsx" / f"beauty_personal_care_top100_fallback_{day}.xlsx"
    build_workbook(output)
    result: dict[str, object] = {"xlsx": str(output), "bytes": output.stat().st_size, "source": "local_sqlite"}
    if not args.no_upload:
        from crawl_amazon_beauty_bestsellers.drive_upload import upload_file

        # A Drive failure must never prevent the local XLSX from being
        # attached and delivered by the enclosing pipeline.
        try:
            result["drive"] = upload_file(output, remote_name=output.name)
        except Exception as exc:
            result["drive_error"] = f"{type(exc).__name__}: {exc}"
    if args.email:
        sys.path.insert(0, str(REPO_ROOT / "scripts"))
        from run_job import _send_email

        link = result.get("drive", {}).get("link", "") if isinstance(result.get("drive"), dict) else ""
        _send_email(
            f"[Amazon Beauty] independent fallback XLSX {time.strftime('%Y-%m-%d %H:%M KST')}",
            "\n".join([
                "Independent fallback workbook completed from local SQLite.",
                "It does not depend on a Google Sheets export.",
                f"Drive: {link}" if link else "Drive upload disabled.",
            ]),
            attachments=[output],
        )
    if args.summary_json:
        args.summary_json.parent.mkdir(parents=True, exist_ok=True)
        args.summary_json.write_text(json.dumps(result, ensure_ascii=False), encoding="utf-8")
    print(result)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
